import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261027   # fixed so the sample is reproducible
N = 50

d = pd.read_csv("scopus_v03.csv", encoding="utf-8-sig")
s = d.sample(n=N, random_state=SEED).reset_index(drop=True)

out = pd.DataFrame({
    "sample_id": [f"T{i:02d}" for i in range(1, N + 1)],
    "year": s["Year"], "title": s["Title"], "source": s["Source title"], "doi": s["DOI"],
    "abstract": s["Abstract"].str.replace(r"\s+", " ", regex=True).str.strip(),
    "my_judgement": "", "my_reason_if_not_relevant": "", "notes": "",
})

with pd.ExcelWriter("precision_sample_v03.xlsx", engine="openpyxl") as w:
    out.to_excel(w, index=False, sheet_name="sample")
    ws = w.sheets["sample"]
    for col, wd in {"A": 9, "B": 7, "C": 42, "D": 26, "E": 24, "F": 85, "G": 16, "H": 30, "I": 30}.items():
        ws.column_dimensions[col].width = wd
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in "GHI":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    dv = DataValidation(type="list", formula1='"relevant,maybe,not relevant"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{N + 1}")

print(len(d), "records;", N, "sampled with seed", SEED)
