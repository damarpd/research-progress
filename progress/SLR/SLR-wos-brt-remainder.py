import re
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261031   # fixed so the order of the sheet is reproducible


def nt(t):
    return re.sub(r"[^a-z0-9]", "", str(t).lower())


def nd(x):
    return x.strip().lower() if isinstance(x, str) and x.strip() else None


# Web of Science records that the broad BRT query returns but v0.3 misses (163), minus
# records already in the merged set (Scopus v0.3 retrieved them) and records already
# in the Scopus broad-BRT missed set (sampled population): what remains was missed by both.
w = pd.read_csv("wos_brt_missed.csv", encoding="utf-8-sig")
m = pd.read_csv("merged_v03.csv", encoding="utf-8-sig")
sm = pd.read_csv("scopus_brt_missed.csv", encoding="utf-8-sig")
w["doi_n"] = w["DOI"].map(nd)
w["tk"] = w["Article Title"].map(nt)
in_merged = w.doi_n.isin(set(m.doi.map(nd).dropna())) | w.tk.isin(set(m.title.map(nt)))
in_scopus_missed = (w.doi_n.isin(set(sm.DOI.map(nd).dropna())) | w.tk.isin(set(sm.Title.map(nt)))) & ~in_merged
rest = w[~in_merged & ~in_scopus_missed].sample(frac=1, random_state=SEED).reset_index(drop=True)

out = pd.DataFrame({
    "sample_id": [f"W{i:02d}" for i in range(1, len(rest) + 1)],
    "year": rest["Publication Year"], "title": rest["Article Title"], "source": rest["Source Title"],
    "doi": rest["DOI"],
    "abstract": rest["Abstract"].str.replace(r"\s+", " ", regex=True).str.strip(),
    "my_judgement": "", "my_reason_if_not_relevant": "", "notes": "",
})

with pd.ExcelWriter("brt_wos_remainder.xlsx", engine="openpyxl") as wr:
    out.to_excel(wr, index=False, sheet_name="sample")
    ws = wr.sheets["sample"]
    for col, wd in {"A": 9, "B": 7, "C": 42, "D": 26, "E": 24, "F": 85, "G": 16, "H": 34, "I": 30}.items():
        ws.column_dimensions[col].width = wd
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in "GHI":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    n = len(out) + 1
    dv = DataValidation(type="list", formula1='"relevant,maybe,not relevant"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{n}")
    dv2 = DataValidation(
        type="list",
        formula1='"BRT planning or operations only,no platform-level boarding/alighting,not transport (other BRT meaning),transit but not BRT platform,other"',
        allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv2)
    dv2.add(f"H2:H{n}")

print(len(w), "WoS records;", in_merged.sum(), "already in merged set;", in_scopus_missed.sum(),
      "in Scopus missed set;", len(rest), "missed by both")
