import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261030   # fixed so the screening order is reproducible

m = pd.read_csv("merged_v03.csv", encoding="utf-8-sig")
order = m.sample(frac=1, random_state=SEED).reset_index(drop=True)   # random order
order.insert(0, "screen_id", [f"S{i:03d}" for i in range(1, len(order) + 1)])

# key linking screening ids back to merged records; keep it out of sight while screening
order[["screen_id", "rec_id", "prior_id", "found_in"]].to_csv("screening_key.csv", index=False)

sheet = pd.DataFrame({
    "screen_id": order["screen_id"], "year": order["year"], "title": order["title"],
    "source": order["source"], "doi": order["doi"],
    "abstract": order["abstract"].str.replace(r"\s+", " ", regex=True).str.strip(),
    "author_keywords": order["author_keywords"],
    "decision": "", "reason_if_exclude": "", "tier_if_include": "", "notes": "",
})

REASONS = ("not boarding/alighting at a platform-vehicle interface,in-vehicle crowding only,"
           "demand or route choice or station location only,operations or timetabling or bunching only,"
           "general pedestrian or crowd flow without platform-level boarding,not transport or other field,other")

with pd.ExcelWriter("screening_sheet.xlsx", engine="openpyxl") as w:
    sheet.to_excel(w, index=False, sheet_name="screening")
    ws = w.sheets["screening"]
    for col, wd in {"A": 10, "B": 7, "C": 44, "D": 26, "E": 24, "F": 90, "G": 30, "H": 12, "I": 38, "J": 14, "K": 30}.items():
        ws.column_dimensions[col].width = wd
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in "HIJK":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    n = len(sheet) + 1
    for col, formula in {"H": '"include,maybe,exclude"', "I": f'"{REASONS}"', "J": '"core,transferable"'}.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")

print(len(sheet), "records in random order, seed", SEED)
